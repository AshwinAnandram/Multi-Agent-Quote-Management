"""Extract reviewable text from common RFQ and quote attachments."""

from io import BytesIO
from pathlib import PurePosixPath


MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 100_000


class AttachmentParseError(ValueError):
    """Raised when an attachment cannot be safely parsed as text."""


def extract_attachment_text(
    filename: str,
    content: bytes,
    mime_type: str | None = None,
) -> str:
    """Extract text from PDF, DOCX, XLSX, or plain-text attachments.

    Scanned PDFs are deliberately reported as unreadable here so the intake
    workflow can hold them for review or OCR instead of silently losing data.
    """
    if not content:
        raise AttachmentParseError("The attachment is empty.")
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise AttachmentParseError("The attachment exceeds the 20 MB limit.")

    suffix = PurePosixPath(filename.lower()).suffix
    if suffix == ".pdf" or mime_type == "application/pdf":
        extracted = _extract_pdf(content)
    elif suffix == ".docx" or mime_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        extracted = _extract_docx(content)
    elif suffix == ".xlsx" or mime_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        extracted = _extract_xlsx(content)
    elif suffix in {".txt", ".csv"} or mime_type in {
        "text/plain",
        "text/csv",
    }:
        extracted = content.decode("utf-8-sig", errors="replace")
    else:
        raise AttachmentParseError(
            f"Unsupported attachment type: {suffix or mime_type or 'unknown'}. "
            "Supported types are PDF, DOCX, XLSX, TXT, and CSV."
        )

    extracted = extracted.strip()
    if not extracted:
        raise AttachmentParseError(
            "No selectable text was found. Scanned PDFs require OCR."
        )
    if len(extracted) > MAX_EXTRACTED_CHARACTERS:
        raise AttachmentParseError(
            "Extracted content exceeds the 100,000 character processing limit."
        )
    return extracted


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content), strict=False)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except AttachmentParseError:
        raise
    except Exception as error:
        raise AttachmentParseError("The PDF could not be read.") from error


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        ]
        return "\n".join(part for part in [*paragraphs, *tables] if part)
    except Exception as error:
        raise AttachmentParseError("The DOCX document could not be read.") from error


def _extract_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
        sections = []
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sections.append(f"[{worksheet.title}]\n" + "\n".join(rows))
        workbook.close()
        return "\n\n".join(sections)
    except Exception as error:
        raise AttachmentParseError("The XLSX workbook could not be read.") from error
